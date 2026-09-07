"""STEP 9 発表準備メモを Word(.docx)として生成する。"""

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from docx.enum.table import WD_TABLE_ALIGNMENT


def set_font(doc, name_ascii="Calibri", name_ja="游ゴシック", size=11):
    style = doc.styles["Normal"]
    style.font.name = name_ascii
    style.font.size = Pt(size)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), name_ja)


def add_heading(doc, text, level):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.name = "Calibri"
        run.element.rPr.rFonts.set(qn("w:eastAsia"), "游ゴシック")
        run.font.color.rgb = RGBColor(0x1F, 0x3B, 0x5C)
    return h


def add_para(doc, text, bold=False, size=11, space_after=6):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = "Calibri"
    run.element.rPr.rFonts.set(qn("w:eastAsia"), "游ゴシック")
    p.paragraph_format.space_after = Pt(space_after)
    return p


def add_code(doc, lines):
    for line in lines:
        p = doc.add_paragraph()
        run = p.add_run(line)
        run.font.name = "Consolas"
        run.element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.left_indent = Pt(18)
    return doc.add_paragraph()


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Light Grid Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h
        cell.paragraphs[0].runs[0].bold = True
    for r, row in enumerate(rows, start=1):
        for c, value in enumerate(row):
            table.rows[r].cells[c].text = value
    doc.add_paragraph()
    return table


doc = Document()
set_font(doc)

# タイトル
title = doc.add_heading("STEP 9 発表準備メモ", level=0)
for run in title.runs:
    run.font.name = "Calibri"
    run.element.rPr.rFonts.set(qn("w:eastAsia"), "游ゴシック")
add_para(doc, "Kaikei Lite(仕訳記録と残高試算表)/ 発表時間: 約 5 分 / 実演は既存の検証用 DB(kaikei-textbook-verify.db、教材の 6 仕訳登録済み)を使用し、データの重複登録は行わない。")

# ── SC-002 と手動確認事項 ─────────────────────────────
add_heading(doc, "SC-002 の内容と手動確認が必要な項目", level=1)
add_para(doc, "SC-002(成功基準): 経験なき経理担当者が、案内だけで最初の仕訳を 1 分以内に登録できる。", bold=True)
add_para(doc, "手動確認が必要な項目(自動テストでは検証できていないもの):")
add_table(
    doc,
    ["項目", "現状", "確認方法"],
    [
        ["SC-002(1 分以内で最初の仕訳登録)", "未検証 — 定量的な UX 指標のため自動テスト対象外にしており、実施していない", "実際に CLI を触ってもらい、かかる時間を計測。発表の実演では追加登録を行わないため、現時点では「未確認」のまま"],
        ["表形式出力の視認性(「人が読める表」)", "自動テストは列構造・JSON 妥当性のみ検証。表示の読みやすさは目視確認していない", "実演時にターミナル表示を目視で確認"],
    ],
)
add_para(doc, "自動検証済みのため手動確認不要なもの: SC-001(不一致 100% 拒否)・SC-003(1,000 件 1 秒以内)・SC-004(取消追跡)・SC-005(全機能 CLI + 両出力)、quickstart 8 シナリオ、教材合格基準 5 項目 — いずれもテストまたは検証スクリプトで確認済み。")

# ── 1. 仕様 ─────────────────────────────
add_heading(doc, "1. 仕様: 逆仕訳の日付・繰越高・金額上限をどう決めたか(約 1 分 20 秒)", level=1)
add_para(doc, "Spec Kit で「仕様 → 明確化 → 計画 → 実装」の順に進めた。特に曖昧だった 3 点を、実装前に質問とチェックリストで決めた。")

add_para(doc, "逆仕訳の日付", bold=True)
add_para(doc, "取消を実行した日(日本時間の当日)を既定値とし、利用者が指定できる。元の仕訳の日付は決して変えない。こうすることで「各仕訳はそれぞれの日付が属する期間の集計に反映される」という規則が破られない。")

add_para(doc, "繰越高", bold=True)
add_para(doc, "試算表の借方合計・貸方合計の列は期間内の仕訳だけで集計し、残高の列には開始日より前の仕訳の累計(繰越高)を含める。期間の両端は含む。実務の「期首からの累計残高」に合わせた。")

add_para(doc, "金額上限", bold=True)
add_para(doc, "当初は「上限なし」で進めていたが、実装計画の段階で「SQLite の整数は 64bit まで」という技術的制約を発見。仕様を勝手に変えず対応案を提示し、「1 明細あたり 9,223,372,036,854,775,807 円以下」に決定。逆に合計・残高の集計には上限を設けないため、集計は Python の多倍長整数で行う方針も同時に確定した。")

# ── 2. 実演 ─────────────────────────────
add_heading(doc, "2. 実演: 貸借不一致の拒否と教材データの試算表(約 2 分)", level=1)
add_para(doc, "検証用 DB(kaikei-textbook-verify.db)には教材の 6 仕訳が登録済み。実演では再登録せず、読み取りと「保存されない失敗登録」のみを行う。")
add_code(doc, [
    "# 前提設定(専用の検証用 DB を使用 — 既存 kaikei.db には触れない)",
    'Set-Location "C:\\Users\\li.zhifeng\\OneDrive - 株式会社テクノスジャパン\\デスクトップ\\cv4-starter\\it-shinjins-week\\it-shinjins-week\\my-project"',
    '$env:PYTHONPATH = "$PWD\\src"',
    '$env:KAIKEI_DB = "$PWD\\kaikei-textbook-verify.db"',
    "",
    "# ① 貸借不一致の仕訳は拒否される(差額 10,000 円を含むエラー)",
    'python -m kaikei entry add --date 2026-04-30 --description "不一致テスト" --debit "現金:300000" --credit "売上高:290000"; echo "exit=$LASTEXITCODE"',
    "# → 一覧を表示すると 6 件のまま(データは一切保存されない)",
    "python -m kaikei entry list",
    "",
    "# ② 教材データ(4/1〜4/30)の残高試算表",
    "#    借方合計=貸方合計 2,462,000、残高合計 1,450,000、現金 188,000",
    "python -m kaikei trial-balance --start 2026-04-01 --end 2026-04-30",
    "",
    "# (時間があれば)現金の総勘定元帳 — 累計残高 188,000 まで追える",
    'python -m kaikei ledger "現金" --start 2026-04-01 --end 2026-04-30',
])
add_para(doc, "ここで伝えること: 憲法の「一致しないデータを保存する経路があってはならない」が、エラー表示 → 一覧確認で目に見える形で動いていること。試算表の数値は手計算した教材の合格基準と一致すること。")

# ── 3. 判定 ─────────────────────────────
add_heading(doc, "3. 判定: 発見・修正した問題と Converged の結果(約 1 分)", level=1)
add_para(doc, "発見した問題 2 件(いずれも実装・検証段階で発見し修正済み):")
add_table(
    doc,
    ["#", "問題", "発見経緯", "対応"],
    [
        ["1", "SQLite の SUM は合計が 64bit を超えるとオーバーフローする", "「各明細は上限以下だが合計が上限を超える仕訳」のテスト", "計画どおり集計を Python 側で行う構造だったため影響はゼロ。テストで「SQL SUM を使わない」ことを明示的に固定した"],
        ["2", "CLI が DB 未初期化のまま動く", "ユニットテストをすり抜け、quickstart の実地検証で発見", "get_connection を自動初期化対応に修正。実地検証の価値を再認識"],
    ],
)
add_para(doc, "収束判定: /speckit-converge の結果、仕様 13 要件・受入シナリオ・Edge Cases 9 件・計画判断 9 件・憲法 5 原則のすべてに対し残差課題ゼロ(Converged)。テスト 80 件すべて合格、教材の合格基準 5 項目も専用 DB 検証で全項目 PASS。")

# ── 4. 学び ─────────────────────────────
add_heading(doc, "4. 学び: 仕様を明確にしてから実装する重要性(約 1 分)", level=1)
add_table(
    doc,
    ["学び", "内容"],
    [
        ["曖昧さは後で必ず高くつく", "逆仕訳の日付・繰越高・残高の表示方法は、曖昧なまま実装していたらテストの期待値そのものが決まらず、後からの手戻りが最大のリスクになっていた。質問とチェックリストで仕様段階で潰せたのが最大の差"],
        ["仕様の曖昧さのテストが実装前のバグを防いだ", "チェックリスト項目「金額の上限はどうするか」がなければ、SQLite の 64bit 制約に実装後に気づき、仕様とコードの双方を後から直すことになっていた"],
        ["ただし仕様だけで完璧ではない", "DB 未初期化のバグは、仕様にもテストにも現れない「実際に動かす」工程で発覚。仕様の明確化と実地検証の両輪が必要"],
        ["残る確認事項", "SC-002(1 分以内で登録できるか)は自動化できておらず未確認。本発表後に、実際に触っていただく形で確認したい"],
    ],
)

doc.save("STEP9_発表準備メモ.docx")
print("saved: STEP9_発表準備メモ.docx")
